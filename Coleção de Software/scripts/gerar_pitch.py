#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera apresentação de pitch — Tupan Water Maker"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_AUTO_SIZE
import os

BASE = "/home/tuliofh01/Documents/Arquivo Acadêmico/PUC-MG/Engenharia de Computação/Disciplinas/2026.2/Iot & PLCs/Tupan Water Maker"
MEDIA = os.path.join(BASE, "Documentação Oficial", "Arquivo de Mídia")
CAD = os.path.join(BASE, "Documentação Oficial", "Arquivos CAD")
OUT = os.path.join(BASE, "Tupan_Water_Maker_Pitch.pptx")

# Cores
AZUL = RGBColor(30, 58, 95)
AZUL_CLARO = RGBColor(220, 234, 247)
AZUL_MEDIO = RGBColor(64, 112, 160)
VERDE = RGBColor(46, 125, 50)
VERDE_CLARO = RGBColor(226, 239, 218)
LARANJA = RGBColor(230, 81, 0)
LARANJA_CLARO = RGBColor(255, 224, 200)
CINZA = RGBColor(89, 89, 89)
CINZA_CLARO = RGBColor(242, 242, 242)
BRANCO = RGBColor(255, 255, 255)
AMARELO = RGBColor(255, 193, 7)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# --- Funções auxiliares ---
def set_bg(slide, color=BRANCO):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color

def add_text(slide, text, x, y, w, h, size=24, color=AZUL, bold=False, align=PP_ALIGN.LEFT, font="Arial", italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return box

def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.6, 0.35, 12.0, 0.7, size=30, color=AZUL, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.62, 1.05, 12.0, 0.4, size=13, color=CINZA, italic=True)
    # linha decorativa
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.45), Inches(1.2), Inches(0.08))
    line.fill.solid(); line.fill.fore_color.rgb = VERDE; line.line.fill.background()

def add_footer(slide, num):
    add_text(slide, f"Tupan Water Maker | PUC-MG Engenharia de Computação", 0.5, 7.15, 10, 0.25, size=8, color=CINZA)
    add_text(slide, str(num), 12.4, 7.15, 0.4, 0.25, size=8, color=CINZA, align=PP_ALIGN.RIGHT)

def add_circle(slide, text, x, y, d, fill_color, text_color=BRANCO, size=18):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color
    tf = shape.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.name = "Arial"; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = text_color
    return shape

def add_rect(slide, text, x, y, w, h, fill_color, text_color=BRANCO, size=14, bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color
    tf = shape.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.name = "Arial"; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = text_color
    return shape

def add_arrow(slide, x1, y1, x2, y2, color=AZUL_MEDIO, width=2):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color; line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line

def add_image(slide, path, x, y, w, h=None):
    if h is None:
        slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    else:
        slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w), height=Inches(h))

# --- Slide 1: Capa ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, AZUL)
# ondas decorativas
for i, (y, c) in enumerate([(5.6, RGBColor(46, 78, 120)), (6.0, RGBColor(46, 125, 50)), (6.4, RGBColor(64, 112, 160))]):
    wave = slide.shapes.add_shape(MSO_SHAPE.ARC, Inches(-1), Inches(y), Inches(15), Inches(2.0))
    wave.fill.solid(); wave.fill.fore_color.rgb = c; wave.fill.transparency = 40; wave.line.fill.background()
add_text(slide, "TUPAN", 0.7, 1.1, 12.0, 1.2, size=64, color=BRANCO, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "WATER MAKER", 0.7, 2.2, 12.0, 0.8, size=34, color=AZUL_CLARO, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "Água do ar. Onde você precisar.", 1.0, 3.3, 11.3, 0.6, size=22, color=BRANCO, align=PP_ALIGN.CENTER)
add_text(slide, "Dispositivo portátil de extração de água atmosférica", 1.0, 4.0, 11.3, 0.4, size=15, color=AZUL_CLARO, align=PP_ALIGN.CENTER)
add_text(slide, "PUC-MG · Engenharia de Computação · IoT & PLCs", 1.0, 5.0, 11.3, 0.35, size=12, color=BRANCO, align=PP_ALIGN.CENTER)
add_text(slide, "Túlio Ferreira Horta", 1.0, 5.5, 11.3, 0.35, size=12, color=BRANCO, align=PP_ALIGN.CENTER)
add_footer(slide, 1)

# --- Slide 2: O problema ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BRANCO)
add_title(slide, "O PROBLEMA", "Escassez hídrica afeta milhões de pessoas")
add_circle(slide, "SECA", 0.9, 2.2, 1.6, LARANJA, size=20)
add_circle(slide, "CUSTO", 0.9, 4.3, 1.6, LARANJA, size=20)
add_circle(slide, "LOGÍSTICA", 0.9, 6.4, 1.6, LARANJA, size=20)
add_arrow(slide, 2.5, 2.6, 4.0, 2.6, LARANJA)
add_arrow(slide, 2.5, 4.7, 4.0, 4.7, LARANJA)
add_arrow(slide, 2.5, 6.8, 4.0, 6.8, LARANJA)
add_rect(slide, "Nordeste brasileiro:\n+ 30 milhões de pessoas\nem situação de risco hídrico", 4.1, 1.9, 3.4, 1.5, LARANJA_CLARO, LARANJA, size=16)
add_rect(slide, "Caminhão-pipa e garrafas:\npreço elevado + emissões", 4.1, 3.7, 3.4, 1.5, LARANJA_CLARO, LARANJA, size=16)
add_rect(slide, "Dependência de infraestrutura\nem áreas remotas", 4.1, 5.5, 3.4, 1.5, LARANJA_CLARO, LARANJA, size=16)
add_rect(slide, "SOLUÇÃO:\nextrair água do ar", 8.0, 3.0, 4.2, 1.6, VERDE_CLARO, VERDE, size=24)
add_footer(slide, 2)

# --- Slide 3: A solução ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BRANCO)
add_title(slide, "A SOLUÇÃO", "Tupan: água potável em qualquer lugar")
add_image(slide, os.path.join(MEDIA, "projSketch_melhorada.jpg"), 0.6, 1.8, 5.6)
add_text(slide, "O Tupan transforma umidade do ar em água potável", 6.7, 2.0, 5.8, 0.8, size=22, color=AZUL, bold=True)
add_rect(slide, "PORTÁTIL", 6.7, 3.1, 2.0, 0.9, AZUL_CLARO, AZUL, size=18)
add_rect(slide, "SUSTENTÁVEL", 8.9, 3.1, 2.4, 0.9, VERDE_CLARO, VERDE, size=18)
add_rect(slide, "CONECTADO", 11.5, 3.1, 1.6, 0.9, LARANJA_CLARO, LARANJA, size=18)
add_text(slide, "• Bateria de carro interna\n• Ligado na tomada\n• Monitoramento via Bluetooth/MQTT\n• Uso exclusivo para beber", 6.7, 4.4, 5.8, 1.6, size=17, color=CINZA)
add_footer(slide, 3)

# --- Slide 4: Como funciona ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BRANCO)
add_title(slide, "COMO FUNCIONA", "Do ar úmido à água potável em 4 etapas")
steps = [
    ("1", "CAPTA", "Ventilador aspira\nar úmido", AZUL_MEDIO),
    ("2", "CONDENSA", "Resfriamento abaixo\ndo ponto de orvalho", AZUL),
    ("3", "FILTRA", "Carvão ativado +\nUV-C", VERDE),
    ("4", "ARMAZENA", "Água pronta\npara beber", LARANJA),
]
for i, (num, title, desc, color) in enumerate(steps):
    x = 0.7 + i * 3.1
    add_circle(slide, num, x, 2.3, 0.9, color, size=22)
    add_text(slide, title, x, 3.35, 2.6, 0.4, size=16, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, desc, x, 3.85, 2.6, 0.8, size=14, color=CINZA, align=PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        add_arrow(slide, x + 1.0, 2.75, x + 2.8, 2.75, CINZA_CLARO, 3)
add_rect(slide, "Ciclo completo em até 8 horas", 3.3, 5.4, 6.7, 0.8, AZUL_CLARO, AZUL, size=20)
add_footer(slide, 4)

# --- Slide 5: Tecnologia ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BRANCO)
add_title(slide, "TECNOLOGIA", "Hardware inteligente + conectividade IoT")
add_image(slide, os.path.join(CAD, "esquema_eletrico.png"), 0.5, 1.8, 6.2)
add_text(slide, "Cérebro eletrônico", 7.2, 1.9, 5.0, 0.5, size=20, color=AZUL, bold=True)
add_rect(slide, "ARDUINO MEGA 2560\nATmega2560 · 16MHz", 7.2, 2.5, 4.6, 0.9, AZUL_CLARO, AZUL, size=16)
add_rect(slide, "BLUETOOTH HC-05\nUART ↔ App Mobile", 7.2, 3.7, 4.6, 0.9, LARANJA_CLARO, LARANJA, size=16)
add_rect(slide, "MQTT\nDashboard em tempo real", 7.2, 4.9, 4.6, 0.9, VERDE_CLARO, VERDE, size=16)
add_text(slide, "Sensores: temperatura · umidade · pressão · nível", 7.2, 6.1, 5.0, 0.35, size=12, color=CINZA, align=PP_ALIGN.CENTER)
add_footer(slide, 5)

# --- Slide 6: Energia ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BRANCO)
add_title(slide, "ENERGIA", "Autonomia para levar água a qualquer lugar")
add_image(slide, os.path.join(MEDIA, "grafico_descarga_bateria.png"), 0.6, 1.8, 6.0)
add_text(slide, "Dual Power", 7.1, 1.9, 5.0, 0.5, size=22, color=AZUL, bold=True)
add_rect(slide, "BATERIA 12V 60Ah\n~20h de operação", 7.1, 2.6, 4.8, 1.0, AZUL_CLARO, AZUL, size=18)
add_rect(slide, "TOMADA 110/220V\nuso contínuo", 7.1, 3.9, 4.8, 1.0, VERDE_CLARO, VERDE, size=18)
add_rect(slide, "MODO ECO\naté 6 dias em standby", 7.1, 5.2, 4.8, 1.0, LARANJA_CLARO, LARANJA, size=18)
add_footer(slide, 6)

# --- Slide 7: Sustentabilidade ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BRANCO)
add_title(slide, "SUSTENTABILIDADE", "Menos desperdício, mais impacto positivo")
add_image(slide, os.path.join(MEDIA, "fluxograma_reciclagem.png"), 0.5, 1.7, 5.4)
add_text(slide, "Impacto positivo", 6.4, 1.9, 5.5, 0.5, size=22, color=VERDE, bold=True)
add_rect(slide, "MATERIAIS\nReciclados + sucata", 6.4, 2.6, 2.6, 1.0, VERDE_CLARO, VERDE, size=16)
add_rect(slide, "ENERGIA\nBaixo consumo", 9.3, 2.6, 2.6, 1.0, VERDE_CLARO, VERDE, size=16)
add_rect(slide, "PRODUTO\nUso para beber", 6.4, 3.9, 2.6, 1.0, VERDE_CLARO, VERDE, size=16)
add_rect(slide, "RECICLAGEM\nReagentes reutilizáveis", 9.3, 3.9, 2.6, 1.0, VERDE_CLARO, VERDE, size=16)
add_text(slide, "Alinhado aos ODS 3, 6 e 12 da ONU", 6.4, 5.5, 5.5, 0.4, size=15, color=CINZA, align=PP_ALIGN.CENTER)
add_footer(slide, 7)

# --- Slide 8: Mercado ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BRANCO)
add_title(slide, "MERCADO", "Onde o Tupan faz diferença")
segments = [
    ("EXPEDIÇÕES", "Trilhas, acampamentos, áreas remotas", AZUL_MEDIO),
    ("COMUNIDADES", "Semiárido, vilarejos, assentamentos", VERDE),
    ("EMERGÊNCIA", "Desastres naturais, crise hídrica", LARANJA),
    ("RESIDENCIAL", "Casas, chácaras, uso diário", AZUL),
]
for i, (title, desc, color) in enumerate(segments):
    col = i % 2
    row = i // 2
    x = 0.8 + col * 6.0
    y = 1.9 + row * 2.4
    add_rect(slide, title, x, y, 2.4, 0.8, color, BRANCO, size=18)
    add_text(slide, desc, x + 2.6, y + 0.1, 3.0, 0.6, size=15, color=CINZA)
add_rect(slide, "1 a 2 pessoas por ciclo", 4.3, 6.2, 4.7, 0.8, AZUL_CLARO, AZUL, size=20)
add_footer(slide, 8)

# --- Slide 9: Modelo de valor ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BRANCO)
add_title(slide, "MODELO DE VALOR", "Por que o Tupan é competitivo")
metrics = [
    ("R$ 180", "custo estimado", LARANJA),
    ("60%", "material reciclado", VERDE),
    ("20h", "autonomia em bateria", AZUL),
    ("0", "emissões locais", AZUL_MEDIO),
]
for i, (val, label, color) in enumerate(metrics):
    x = 0.8 + i * 3.1
    add_text(slide, val, x, 2.2, 2.8, 0.8, size=32, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x, 3.1, 2.8, 0.4, size=14, color=CINZA, align=PP_ALIGN.CENTER)
add_text(slide, "Água potável onde não há infraestrutura", 2.0, 4.6, 9.3, 0.6, size=24, color=AZUL, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "Baixo custo · Alta mobilidade · Tecnologia acessível", 2.0, 5.3, 9.3, 0.4, size=15, color=CINZA, align=PP_ALIGN.CENTER)
add_footer(slide, 9)

# --- Slide 10: Próximos passos ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, AZUL)
add_text(slide, "PRÓXIMOS PASSOS", 0.7, 0.6, 12.0, 0.8, size=32, color=BRANCO, bold=True, align=PP_ALIGN.CENTER)
steps = [
    ("1", "Protótipo funcional", "montagem e testes"),
    ("2", "Validação de qualidade", "água potável"),
    ("3", "Piloto em comunidade", "campo real"),
    ("4", "Escala sustentável", "produção local"),
]
for i, (num, title, desc) in enumerate(steps):
    x = 0.9 + i * 3.1
    add_circle(slide, num, x, 2.2, 0.9, VERDE, size=22)
    add_text(slide, title, x, 3.3, 2.7, 0.4, size=16, color=BRANCO, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, desc, x, 3.8, 2.7, 0.35, size=13, color=AZUL_CLARO, align=PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        add_arrow(slide, x + 1.0, 2.65, x + 2.8, 2.65, RGBColor(120, 160, 200), 3)
add_text(slide, "Tupan Water Maker", 1.0, 5.6, 11.3, 0.5, size=22, color=BRANCO, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "Água do ar. Onde você precisar.", 1.0, 6.2, 11.3, 0.4, size=16, color=AZUL_CLARO, align=PP_ALIGN.CENTER)
add_footer(slide, 10)

# Salvar
prs.save(OUT)
print("Apresentação salva:", OUT)
print("Slides:", len(prs.slides))
